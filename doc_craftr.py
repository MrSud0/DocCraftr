import os
import argparse
import random
from fpdf import FPDF
from docx import Document
import csv
import json
import xlwt

# Dummy text for all files
dummy_text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua."

# Pool of 50 random file names closer to a corporate environment
file_name_pool = [
    "annual_report", "budget_summary", "client_list", "project_plan", "meeting_minutes", "financial_statement", "hr_policies", "company_profile", "team_structure", "roadmap_2024",
    "employee_directory", "quarterly_review", "contract_template", "business_strategy", "training_materials", "marketing_plan", "vendor_contacts", "invoice_template", "audit_log", "sales_forecast",
    "action_items", "service_agreement", "customer_feedback", "risk_assessment", "incident_report", "product_catalog", "board_meeting_agenda", "expense_report", "security_guidelines", "new_hire_onboarding",
    "performance_review", "supply_chain_overview", "market_analysis", "internal_memo", "brand_guidelines", "data_privacy", "nda_agreement", "client_proposal", "system_architecture", "dev_roadmap",
    "asset_inventory", "operational_plan", "monthly_expenses", "growth_strategy", "it_policies", "office_layout", "legal_notice", "service_manual", "training_schedule", "contractor_info"
]

# Function to generate .txt files
def generate_txt_file(path):
    with open(path, 'w') as file:
        file.write(dummy_text)

# Function to generate .pdf files
def generate_pdf_file(path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', size=12)
    pdf.multi_cell(200, 10, txt=dummy_text)
    pdf.output(path)

# Function to generate .docx files
def generate_docx_file(path):
    document = Document()
    document.add_paragraph(dummy_text)
    document.save(path)

# Function to generate .csv files
def generate_csv_file(path):
    with open(path, mode='w', newline='') as file:
        writer = csv.writer(file)
        writer.writerow(["Column1", "Column2", "Column3"])
        writer.writerow(["Data1", "Data2", "Data3"])
        writer.writerow(["Data4", "Data5", "Data6"])

# Function to generate .json files
def generate_json_file(path):
    data = {
        "key1": "value1",
        "key2": "value2",
        "key3": "value3"
    }
    with open(path, 'w') as file:
        json.dump(data, file, indent=4)

# Function to generate .xls files
def generate_xls_file(path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet1")
    sheet.write(0, 0, "Header1")
    sheet.write(0, 1, "Header2")
    sheet.write(0, 2, "Header3")
    sheet.write(1, 0, "Value1")
    sheet.write(1, 1, "Value2")
    sheet.write(1, 2, "Value3")
    workbook.save(path)

# Function to generate dummy files based on specified format
def generate_files(num_files, folder, formats):
    used_file_names = set()

    for i in range(num_files):
        file_format = random.choice(formats)
        file_name = random.choice(file_name_pool)

        # Ensure unique file names
        while file_name in used_file_names:
            file_name = random.choice(file_name_pool)
        used_file_names.add(file_name)

        filename = f"{file_name}.{file_format}"
        file_path = os.path.join(folder, filename)

        if file_format == 'txt':
            generate_txt_file(file_path)
        elif file_format == 'pdf':
            generate_pdf_file(file_path)
        elif file_format == 'docx':
            generate_docx_file(file_path)
        elif file_format == 'csv':
            generate_csv_file(file_path)
        elif file_format == 'json':
            generate_json_file(file_path)
        elif file_format == 'xls':
            generate_xls_file(file_path)
        else:
            print(f"Unsupported file format: {file_format}")

# Function to spread generated files into random subdirectories
def spread_files(root_folder):
    all_subdirs = []
    for root, dirs, _ in os.walk(root_folder):
        for d in dirs:
            all_subdirs.append(os.path.join(root, d))

    # If no subdirectories found, skip spreading
    if not all_subdirs:
        print("No subdirectories found to spread files into.")
        return

    # Move each file in the root folder to a random subdirectory
    for filename in os.listdir(root_folder):
        file_path = os.path.join(root_folder, filename)
        if os.path.isfile(file_path):
            target_dir = random.choice(all_subdirs)
            target_path = os.path.join(target_dir, filename)
            os.rename(file_path, target_path)
            print(f"Moved {filename} to {target_dir}")

# Main function to parse CLI arguments and generate files
def main():
    parser = argparse.ArgumentParser(description="Generate a mix of .txt, .pdf, .docx, and other documents with dummy data.")
    parser.add_argument('-n', '--number', type=int, required=True, help="Number of files to generate")
    parser.add_argument('-f', '--folder', type=str, required=True, help="Root path of the destination folder")
    parser.add_argument('-m', '--mix', type=str, default='txt,pdf,docx', help="Comma separated list of formats to generate (e.g., txt,pdf,docx,csv,json,xls)")
    parser.add_argument('-s', '--spread', action='store_true', help="Enable spreading of files into random subdirectories")

    args = parser.parse_args()
    num_files = args.number
    root_folder = args.folder
    formats = args.mix.split(',')

    # Ensure the folder exists
    if not os.path.exists(root_folder):
        os.makedirs(root_folder)

    # Generate files
    generate_files(num_files, root_folder, formats)
    print(f"{num_files} files generated in {root_folder} with formats: {', '.join(formats)}")

    # Spread files to random subdirectories if --spread is enabled
    if args.spread:
        spread_files(root_folder)

if __name__ == "__main__":
    main()
